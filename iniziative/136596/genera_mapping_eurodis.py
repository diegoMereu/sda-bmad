import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ICD_PREADVISING_LDV_EVENT_SOURCE_v11.xlsx")

# Categorie
D, T, C, G, F1, CE, AV, NA, DD, NV = (
    "DIRETTO", "TRASFORMATO", "COSTANTE", "GENERATO", "F1", "CE-RULES",
    "A VALLE", "N/A", "DA DEFINIRE", "NON VALORIZZATO")

S = "spedizione."
M, DE = S + "mittente.", S + "destinatario."

def addr(prefix, src, extra_note=""):
    """Regole comuni indirizzo/contatti (mittente e destinatario)."""
    return {
        prefix + "ragione_sociale": (D, f"{src}.companyName", "40 -> 40", ""),
        prefix + "rif_contatto": (D, f"{src}.contactName", "Valore integrale (sorgente 40, ICD 35): non si tronca", "PA-16"),
        prefix + "email": (D, f"{src}.contactEMail", "Valore integrale (sorgente 80, ICD 50): non si tronca", "PA-16"),
        prefix + "cellulare": (D, f"{src}.contactSMS", "Valore integrale, formato +XX... (sorgente 35, ICD 30): non si tronca", "PA-16"),
        prefix + "telefono": (D, f"{src}.contactTelephone", "30 -> 30", ""),
        prefix + "indirizzo": (D, f"{src}.street", "40 -> 40", ""),
        prefix + "numero_civico": (D, f"{src}.streetNumber", "Valore integrale (sorgente 40, ICD 5): non si tronca", "PA-16"),
        prefix + "cap": (D, f"{src}.zipcode", "9 -> 9. CAP esteri ammessi", ""),
        prefix + "localita": (D, f"{src}.city", "35 -> 40", ""),
        prefix + "provincia": (DD, f"{src}.region", "region (40, testo libero) non e' la sigla provincia: valorizzare solo se IT e sigla valida, altrimenti lasciare alla normalizzazione", "PA-22"),
        prefix + "codice_nazione_iso_2": (D, f"{src}.countryCode", "ISO 3166 alpha-2", ""),
        prefix + "codice_nazione": (T, f"{src}.countryCode", "Decodifica ISO2 -> ISO3", ""),
        prefix + "nazione": (T, f"{src}.countryCode", "Decodifica ISO2 -> descrizione", ""),
        prefix + "codice_nazione_poste": (T, f"{src}.countryCode", "Decodifica ISO2 -> codice Poste (ISO4), tabella esistente", ""),
        prefix + "note1": (D, f"{src}.additionalAddressInfo1", "Valore integrale (sorgente 40): non si tronca", "PA-16"),
        prefix + "note2": (D, f"{src}.additionalAddressInfo2", "Valore integrale (sorgente 40): non si tronca", "PA-16"),
    }

R = {
    # --- testata ---
    "progressivo": (NA, "", "Specifico dei CSV PDB", ""),
    "upload_time": (G, "", "Istante di ricezione della risposta /shipments/v1", ""),
    "dataInizioFlusso": (G, "", "Inizio esecuzione batch (UTC)", ""),
    "instantInizioFlusso": (G, "", "Inizio esecuzione batch (ISO 8601 UTC)", ""),
    "link_file_fisico_inviato": (G, "", "Path del payload grezzo salvato (storico, PA-02), se su blob storage", "PA-02"),
    "errori": (AV, "", "Valorizzato dalla validazione", "PA-08"),
    "versione_tracciato": (C, "", "1 (ICD v11)", ""),
    "canale": (C, "", '"EURODIS" (gia previsto fra i valori ICD)', ""),
    "id_canale": (DD, "", "Codice canale a 3 caratteri per EURODIS", "PA-15"),
    "id_preadvising": (G, "", "UUID.randomUUID().toString() (decisione 22-Oct). Attenzione: 36 caratteri, ICD dichiara 15", "PA-15"),
    "flag_autoproduttore": (C, "", '"S"', ""),
    "id_file": (G, "", "Id esecuzione batch (una chiamata = un 'file')", "PA-15"),
    # --- cliente: EURODIS come cliente SDA ---
    "cliente.no_firma": (DD, "deliveryService.deliveryType?", "Da capire se un deliveryType EURODIS implica consegna senza firma", ""),
    # --- spedizione ---
    S + "numero_colli": (D, "numberOfHandlingUnits", "Verificare coerenza con len(collis)", ""),
    S + "rif_spedizione": (T, "edReference", "Proposta: numero EURODIS, chiave per il ritorno tracking (13 cifre <= 25)", "PA-11"),
    S + "id_ldv_partner_master": (D, "dispatchPartnersDomesticShipmentNumber", "Numero spedizione del partner mittente (35 <= 50)", "PA-11"),
    S + "id_partner": (C, "", '"EURODIS" o codice consorzio', "PA-15"),
    S + "descrizione_contenuto": (D, "customsService.descriptionOfGoods", "Solo se presente customsService (20 <= 30)", ""),
    S + "valore": (T, "customsService.commercialValueOfGoods", "x100 (centesimi), solo se currency = EUR", "PA-22"),
    S + "codice_tipo_spedizione": (DD, "collis[].identifierPalletOrCollo?", 'Valori ICD "P","C": verificare il significato', ""),
    S + "codice_servizio_pagamento": (DD, "termsOfDeliveryText?", 'Valori ICD "N","P","T"', ""),
    S + "tipo_spedizione": (T, "shipperAddress.countryCode", "Internazionale se mittente non IT", ""),
    S + "verso": (DD, "shipmentTypeIndicator", "Resi (R) e rettifiche: come si rappresentano", "PA-17"),
    S + "data_ordine": (T, "clearanceTimestamp", "yyyy-MM-dd HH:mm:ss", ""),
    S + "data_spedizione": (T, "clearanceTimestamp", "Formato 8 caratteri (yyyyMMdd?). In alternativa entranceTimestamp (presa fisica)", ""),
    S + "ora_spedizione": (T, "clearanceTimestamp", "HH:mm:ss", ""),
    S + "flag_stampa_ldv": (C, "", '"N": etichetta prodotta da EURODIS (autoproduttore)', ""),
    S + "codice_prodotto": (CE, "shipmentTypeIndicator, lineHaulIndicator, deliveryService.deliveryService, deliveryService.deliveryType", "Da ce-rules (PA-03)", "PA-03"),
    S + "codici_servizi": (CE, "deliveryService.deliveryService, deliveryService.deliveryType, deliveryService.deliveryBookInRequired, cashOnDeliveryService, customsService, pudoID", "Da ce-rules", "PA-03"),
    S + "servizi_spedizione": (CE, "", "Da ce-rules", "PA-03"),
    S + "codice_servizio": (CE, "", "Da ce-rules", "PA-03"),
    S + "id_servizio": (CE, "", "Da ce-rules", "PA-03"),
    S + "id_accessori_base": (CE, "", "Da ce-rules", "PA-03"),
    S + "id_sotto_accessori": (CE, "", "Da ce-rules", "PA-03"),
    S + "codice_prodotto_aptus": (CE, "", "Da ce-rules", "PA-03"),
    S + "notifiche_sms": (T, "consigneePreAdvice.smsPreAdvice", 'true -> "S", false -> "N". Semantica campo ICD da confermare', "PA-18"),
    S + "notifiche_email": (T, "consigneePreAdvice.eMailPreAdvice", 'true -> "S", false -> "N". Semantica campo ICD da confermare', "PA-18"),
    S + "importi.nazione_mittente_iso_cc": (D, "shipperAddress.countryCode", "", ""),
    S + "importi.nazione_mittente_descrizione": (T, "shipperAddress.countryCode", "Decodifica ISO2 -> descrizione", ""),
    S + "importi.status_doganale": (DD, "customsService", "Spedizioni extra-UE: perimetro?", "PA-13"),
    # --- destinatario extra ---
    DE + "prefisso_cellulare": (T, "consigneeAddress.contactSMS", "Prefisso +XX estratto dal numero", ""),
    # --- internazionale ---
    S + "internazionale.id_ldv_itz": (DD, "edReference / europlateBarcode", 'ICD: "previsto per i casi ERS e Eurodis". Candidato naturale per la LDV EURODIS', "PA-11"),
    S + "internazionale.descrizione": (D, "customsService.descriptionOfGoods", "", ""),
    S + "internazionale.peso": (T, "shipmentWeight", "kg -> grammi (x1000)", ""),
    S + "internazionale.valuta": (D, "customsService.currency", "", ""),
    S + "internazionale.valore_spedizione": (D, "customsService.commercialValueOfGoods", "In euro", "PA-22"),
    S + "internazionale.addebito_oneri": (T, "termsOfDeliveryText", 'Proposta: DDC/DDX -> "M", DPU/EXW -> "D". Da confermare', "PA-13"),
    S + "internazionale.vettore": (C, "", '"EURODIS"', ""),
    S + "internazionale.nazionalita": (D, "shipperAddress.countryCode", "", ""),
    # --- assicurata: EURODIS non la prevede ---
    S + "assicurata.spedizione_assicurata": (C, "", '"N": EURODIS non trasmette assicurazione', ""),
    # --- contrassegno ---
    S + "contrassegno.flag": (T, "cashOnDeliveryService", '"S" se presente, altrimenti "N"', ""),
    S + "contrassegno.importo_contrassegno": (D, "cashOnDeliveryService.amount", "Se currency <> EUR: scarto?", "PA-19"),
    S + "contrassegno.tipo_pagamento_contrassegno": (DD, "cashOnDeliveryService.identifierForCODCollection", "3 = contanti, 5 = assegno -> codici ICD (CON? ACM/ABM?)", "PA-19"),
    # --- consegna ---
    S + "consegna.data_consegna": (T, "deliveryService.deliveryDateWindowStart", "yyyy-MM-dd -> 8 caratteri", ""),
    S + "consegna.inizio_ora_consegna": (T, "deliveryService.deliveryTimeWindowStart", "HH:mm:ss -> HH:mm", ""),
    S + "consegna.fine_ora_consegna": (T, "deliveryService.deliveryTimeWindowEnd", "HH:mm:ss -> HH:mm", ""),
    S + "consegna.note": (D, "consigneeAddress.additionalConsigneeInfo", "70 caratteri, nessun limite ICD", ""),
    S + "consegna.modalita": (DD, "deliveryService.deliveryType", "Decodifica S/N/D/I/B/C -> modalita SDA", "PA-17"),
    S + "consegna.codice_servizio_accessorio": (CE, "", "Da ce-rules", "PA-03"),
    S + "consegna.servizi_spedizione": (CE, "", "Da ce-rules", "PA-03"),
    S + "consegna.mancata_consegna": (DD, "", 'Nessuna sorgente EURODIS: default "R" (reso)?', "PA-17"),
    # --- collo ---
    S + "collo.id_ldv": (DD, "", "Dipende dalla decisione LDV", "PA-11"),
    S + "collo.id_ldv_madre": (DD, "", "Multicollo: dipende dalla decisione LDV", "PA-11"),
    S + "collo.id_ldv_figlio": (DD, "", "Multicollo: dipende dalla decisione LDV", "PA-11"),
    S + "collo.id_collo": (D, "collis[].europlateBarcode", "24 <= 25", "PA-11"),
    S + "collo.tipo_collo": (T, "collis[] (posizione)", "Solo multicollo: madre/figlio", "PA-11"),
    S + "collo.numero_colli_parz": (T, "collis[] (posizione)", "Progressivo collo", ""),
    S + "collo.peso": (D, "collis[].weightOfUnit", "kg", ""),
    S + "collo.lunghezza": (D, "collis[].length", "cm", ""),
    S + "collo.altezza": (D, "collis[].height", "cm", ""),
    S + "collo.profondita": (D, "collis[].width", "cm", ""),
    S + "collo.codice_imballo": (T, "collis[].identifierPalletOrCollo", 'Proposta: P/P+ -> "PL", C/C+ -> "CL"', ""),
    S + "collo.nazionalita": (D, "shipperAddress.countryCode", "", ""),
    S + "collo.articolo.codice_taric": (T, "customsService.dutyCode", "Un solo articolo per spedizione; valore integrale (sorgente 12, ICD 10)", "PA-13"),
    S + "collo.articolo.english_descr": (D, "customsService.descriptionOfGoods", "", "PA-13"),
    S + "collo.articolo.valore": (T, "customsService.commercialValueOfGoods", "x100 (centesimi)", "PA-13"),
    S + "collo.articolo.valuta": (D, "customsService.currency", "", "PA-13"),
}
R.update(addr(M, "shipperAddress"))
R.update(addr(DE, "consigneeAddress"))
R[DE + "cap"] = (D, "consigneeAddress.zipcode", "Obbligatorio ICD; obbligatorio anche in EURODIS", "")
R[DE + "localita"] = (D, "consigneeAddress.city", "Obbligatorio ICD; obbligatorio anche in EURODIS", "")

# Allineamento all'esempio del flusso esistente EUD_INVIOFILE (LDV prodotta da SDA),
# in esempio-eud-inviofile.json. Il flusso autoproduttore deve differire di poco.
EX = "nell'esempio EUD_INVIOFILE"
for p, src in ((M, "shipperAddress"), (DE, "consigneeAddress")):
    R[p + "codice_nazione"] = (D, f"{src}.countryCode", f'ISO2 cosi com\'e\' ("AT", "IT" {EX}), anche se l\'ICD dice ISO3', "")
    R[p + "codice_nazione_iso_2"] = (NV, "", f"Vuoto {EX}", "")
    R[p + "codice_nazione_poste"] = (NV, "", f"Vuoto {EX}", "")
R.update({
    "canale": (DD, "", f'"EURODIS" (valore ICD) o un nuovo valore sul modello di "EUD_INVIOFILE" (flusso esistente)?', "PA-15"),
    "id_file": (G, "", f"UUID dell'esecuzione batch (una chiamata = un 'file'); UUID anche {EX}", ""),
    "cliente.codice_login_cliente": (F1, "", f"Codice cliente EURODIS ({EX} uguale a id_cliente e codice_auc)", "PA-10"),
    "cliente.id_cliente": (F1, "", "Codice cliente EURODIS", "PA-10"),
    "cliente.codice_auc": (F1, "", "Codice cliente EURODIS", "PA-10"),
    "cliente.codice_contratto": (F1, "", "Contratto SFA di EURODIS (SFA-...)", "PA-10"),
    "cliente.ragione_sociale": (D, "shipperAddress.companyName", f"{EX} coincide con la ragione sociale del mittente", ""),
    S + "rif_spedizione": (T, "edReference", "Decisione 22-Oct. Numero decimale (max 13 cifre) -> stringa senza decimali (ICD 25)", "PA-11"),
    S + "data_spedizione": (T, "clearanceTimestamp", f"yyyy-MM-dd (come {EX}). In alternativa entranceTimestamp (presa fisica)", ""),
    S + "flag_stampa_ldv": (C, "", f'"N": etichetta prodotta da EURODIS (autoproduttore). Vuoto {EX}', ""),
    S + "formato_stampa_ldv": (NV, "", f'"A4" {EX}; non serve se la LDV la produce EURODIS', ""),
    S + "flag_attivazione_ritiro": (C, "", f'"N" (come {EX})', ""),
    S + "rif_prodotto_cliente": (DD, "", f'"APT000902" {EX}: costante di contratto o derivato?', "PA-23"),
    S + "centro_costo": (DD, "", f'"CDC-00000014" {EX}: costante di contratto (F1)?', "PA-23"),
    S + "assicurata.spedizione_assicurata": (C, "", f"Vuoto, importo 0 (come {EX}): EURODIS non trasmette assicurazione", ""),
    # PA-11, decisione 22-Oct: consignorsColloBarcode = LDV del collo, europlateBarcode = id_ldv_itz
    S + "collo.id_ldv": (D, "collis[].consignorsColloBarcode", "LDV del collo (decisione 22-Oct). Vuoto in EUD_INVIOFILE perche' li' la genera SDA", "PA-11"),
    S + "internazionale.id_ldv_itz": (D, "collis[].europlateBarcode", "Decisione 22-Oct. ICD: \"previsto per i casi ERS e Eurodis\". Campo a livello spedizione, barcode a livello collo: con piu' colli quale si usa?", "PA-11"),
    S + "collo.id_collo": (DD, "", f'"C-MR-C012-15702" {EX}: da quale campo arriva?', "PA-11"),
    S + "collo.peso": (T, "collis[].weightOfUnit", f"kg -> grammi (x1000): peso 2000 {EX}", ""),
})

# Decisioni del 23-Oct dai commenti di Diego (colonna "Commenti" del foglio Mapping inverso).
# I dubbi funzionali aperti sono in punti-aperti-analisi-2026-10-23.md (DF-nn).
K = "commento Diego 23-Oct"
for p in (M, DE):
    R[p + "codice_nazione_iso_2"] = (D, ("shipperAddress" if p == M else "consigneeAddress") + ".countryCode",
                                     f"countryCode e' ISO2: va solo qui ({K})", "")
    for f in ("codice_nazione", "nazione", "codice_nazione_poste"):
        R[p + f] = (NV, "", f"countryCode va solo in codice_nazione_iso_2 ({K})", "DF-12")
for f in ("importi.nazione_mittente_iso_cc", "importi.nazione_mittente_descrizione",
          "collo.nazionalita", "internazionale.nazionalita"):
    R[S + f] = (NV, "", f"countryCode va solo in codice_nazione_iso_2 ({K})", "")
for f, why in (("verso", "shipmentTypeIndicator da ignorare"), ("tipo_spedizione", "non deriva dal paese del mittente"),
               ("data_ordine", "clearanceTimestamp da ignorare"), ("codice_servizio_pagamento", "termsOfDeliveryText da non mappare"),
               ("internazionale.addebito_oneri", "termsOfDeliveryText da non mappare")):
    R[S + f] = (NV, "", f"{why} ({K})", "")
# Prodotto e servizi sono interni SDA: l'applicativo li recupera chiamando ce-rules con id_ldv
for p, r in list(R.items()):
    if r[0] == CE:
        R[p] = (CE, "", f"Interno SDA: ce-rules chiamato con id_ldv in ingresso ({K})", "PA-03")
R[S + "codici_servizi"] = (CE, "", f"Da ce-rules con id_ldv. Aperto: come si riflettono servizio di consegna EURODIS e contrassegno ({K})", "DF-05, DF-07")
R.update({
    S + "data_spedizione": (T, "entranceTimestamp", f"yyyy-MM-dd ({K}: clearanceTimestamp si ignora)", "DF-01"),
    S + "ora_spedizione": (T, "entranceTimestamp", f"HH:mm:ss ({K})", "DF-01"),
    DE + "prefisso_cellulare": (NV, "", f"contactSMS va tutto in cellulare, prefisso compreso ({K})", ""),
    S + "consegna.note": (T, "consigneeAddress.additionalConsigneeInfo, shipperAddress.additionalConsigneeInfo",
                          f"Istruzioni di consegna: concatenare i due campi se entrambi presenti ({K})", ""),
    "cliente.no_firma": (T, "deliveryService.deliveryType", f'S (firma richiesta) -> "N"; N (firma non richiesta) -> "S"; altri valori -> vuoto ({K})', ""),
    S + "consegna.data_consegna": (T, "deliveryService.deliveryDateWindowStart", f"yyyy-MM-dd ({K})", "DF-06"),
    S + "assicurata.importo_assicurata": (D, "customsService.commercialValueOfGoods", f"Valore merce 'for insurance reasons' secondo lo swagger ({K})", "DF-08"),
    S + "assicurata.spedizione_assicurata": (DD, "customsService.commercialValueOfGoods", f'"S" se il valore merce e\' presente? ({K})', "DF-08"),
    S + "id_ldv_associata": (D, "collis[].referencedEuroplateBarcode", f"Solo resi. ICD: campo di spedizione di 3 caratteri, il barcode e' per collo e di 24 ({K})", "DF-10"),
    S + "notifiche_sms": (T, "consigneePreAdvice.smsPreAdvice", f'true -> "S", false -> "N" ({K}: mappato cosi, da confermare)', "DF-09"),
    S + "notifiche_email": (T, "consigneePreAdvice.eMailPreAdvice", f'true -> "S", false -> "N" ({K}: mappato cosi, da confermare)', "DF-09"),
})
R["progressivo"] = (G, "", "Progressivo 1..n della spedizione nell'esecuzione batch: obbligatorio per BEP (fileRecordOrder). Da confermare", "")

# Servizi accessori EURODIS -> formato PDB (proposta del 29-Oct, accessori-eurodis-pdb-proposta.md).
# Fonti: riscontro SDA/Chiapponi 11-set (flusso via file: ipotesi per l'API), FlussiPreadvisingInfo.csv,
# ce-dichiarato-ldv-pdb-sync. Il ms sync legge i codici APT da codici_servizi e le sottochiavi dai campi ICD.
P = "proposta 29-Oct"
R.update({
    S + "codici_servizi": (T, "deliveryService.deliveryService, cashOnDeliveryService, consigneeAddress.pudoID",
                           f'Lista APT separata da pipe ({P}): 6 -> APT000910, E -> APT000911, F -> APT000906, M -> APT000907, '
                           'B -> APT000908, W -> APT000908 (matrice di test SDA, da confermare), Z + pudoID -> APT000947, cashOnDeliveryService presente -> APT000918. '
                           'APT000945 (multicollo) NON si scrive: lo aggiunge il ms sync. A = spedizione normale. '
                           'Valori D, 7, T non coperti (DF-21); fuori perimetro: assicurazione, Locker, Uffici Postali', "DF-05, DF-07, DF-21, DF-22"),
    S + "consegna.codice_frazionario": (T, "consigneeAddress.pudoID",
                                        f"Solo con deliveryService = Z: PUDO ID, e' il node di APT000947 ({P})", "DF-14, DF-23"),
    S + "consegna.denominazione_frazionario": (DD, "",
                                               f"Nome del Punto Poste: e' il name di APT000947, obbligatorio non vuoto. Nel flusso via file il nome e' in RT_30_030, opzionale e non garantito (RT_30_080 e' il destinatario effettivo): cosa fare se manca ({P})", "DF-23"),
    S + "contrassegno.importo_contrassegno": (D, "cashOnDeliveryService.amount",
                                              f"2 decimali, maggiore di 0, max 999999.99 (CSV accessori). Se currency <> EUR: scarto? ({P})", "DF-24"),
    S + "contrassegno.tipo_pagamento_contrassegno": (T, "cashOnDeliveryService.identifierForCODCollection",
                                                     f"3 -> CON, 5 -> ABM (riscontro SDA, da confermare; mai ACM). Non ammesso con Punto Poste ({P})", "DF-24"),
    S + "assicurata.spedizione_assicurata": (NV, "", f"Assicurazione fuori perimetro (riscontro SDA): non valorizzare ({P})", "DF-08"),
    S + "assicurata.importo_assicurata": (NV, "", f"Assicurazione fuori perimetro (riscontro SDA): non usare commercialValueOfGoods ({P})", "DF-08"),
    S + "assicurata.flag": (NV, "", f"Assicurazione fuori perimetro (riscontro SDA) ({P})", "DF-08"),
    S + "consegna.mancata_consegna": (DD, "", f"Reso = rientro al mittente. APT000930 non mappato finche' non si decide se serve e che valore ha return (A o R) ({P})", "DF-25"),
    S + "consegna.data_consegna": (T, "deliveryService.deliveryDateWindowStart", f"yyyy-MM-dd ({K}). Il servizio T non ha APT (DF-21) ({P})", "DF-06, DF-21"),
    S + "notifiche_sms": (DD, "consigneePreAdvice.smsPreAdvice", f"SDA notifica da sola se ci sono i contatti e non usa RT_65 (riscontro SDA): valutare se non mapparlo ({P})", "DF-09"),
    S + "notifiche_email": (DD, "consigneePreAdvice.eMailPreAdvice", f"SDA notifica da sola se ci sono i contatti e non usa RT_65 (riscontro SDA): valutare se non mapparlo ({P})", "DF-09"),
    S + "rif_prodotto_cliente": (DD, "", f'Il ms sync usa questo campo come prodotto e, se vuoto, il valore di config (APT000901, PDB Express) ({P})', "DF-11, PA-23"),
    S + "centro_costo": (DD, "", f"Il ms sync, se vuoto, usa il centro di costo di config; un solo centro statico fornito da Poste (riscontro SDA) ({P})", "DF-11, PA-23"),
    S + "data_spedizione": (T, "entranceTimestamp", f"yyyy-MM-dd ({K}: clearanceTimestamp si ignora). Conflitto: il riscontro SDA mappa clearanceTimestamp -> shipmentDate ({P})", "DF-01"),
})
# Il ms sync legge il paese BEP da codice_nazione e non risolve l'ISO4 (AnagraficaClient mai chiamato): vedi DF-12 e DF-27
for p in (M, DE):
    R[p + "codice_nazione"] = (NV, "", f"countryCode va solo in codice_nazione_iso_2 ({K}). Attenzione: il ms sync legge il paese BEP da qui e non risolve ISO4; destinatario = ITA1 ({P})", "DF-12, DF-27")
GAP_23 = {
    "shipmentTypeIndicator": ("", "", f"Da ignorare ({K})", ""),
    "lineHaulIndicator": ("", "", f"Da ignorare ({K})", ""),
    "clearanceTimestamp": ("", "", f"Da ignorare: per data e ora spedizione si usa entranceTimestamp ({K})", ""),
    "truckLicensePlate": ("", "", f"Non interessa ({K})", ""),
    "termsOfDeliveryText": ("", "", f"Da non mappare ({K})", ""),
    "consigneeAddress.gatewayAssigned": ("", "", f"Da non mappare ({K})", ""),
    "shipperAddress.gatewayAssigned": ("", "", f"Da non mappare ({K})", ""),
    "consigneeAddress.pudoID": ("", "", f"Da non mappare per ora ({K})", "DF-14"),
    "consigneeAddress.alternativeContactTelephone": ("", "", f"Il tracciato non ha un secondo numero di telefono: da chiedere se serve ({K})", "DF-15"),
    "shipperAddress.alternativeContactTelephone": ("", "", f"Il tracciato non ha un secondo numero di telefono: da chiedere se serve ({K})", "DF-15"),
    "deliveryService.deliveryService": ("", "", f"Come si traduce in servizi SDA ({K})", "DF-05"),
}

def default(path, name):
    if "normalizzat" in name or name.endswith("_tipo") and name.split("_")[0] in ("cellulare", "telefono") \
            or ".linee_alternative" in path or ".frazionario_inesito" in path \
            or name in ("cap_fittizio", "flag_cap_fittizio", "esito_normalizzazione", "descrizione_normalizzazione"):
        return (AV, "", "Arricchimento / normalizzazione a valle del topic pre-enrichment", "")
    if path.startswith("ritiro") or path.startswith(S + "pickup"):
        return (NA, "", "EURODIS consegna in hub SDA: niente ritiro/pickup in questo flusso", "")
    if path.startswith("cliente"):
        return (NV, "", "Vuoto o assente nell'esempio EUD_INVIOFILE", "")
    if path.startswith(S + "consegna.") and name in ("filiale", "filiale_ritiro", "id_filiale_stampa", "codice_frazionario",
                                                       "denominazione_frazionario", "codice_giro", "id_fil_dest", "latitudine", "longitudine"):
        return (AV, "", "Calcolato a valle", "")
    if name.startswith("dhl_") or name in ("id_ldv_ups", "id_ldv_tt", "consegnabile_ap_ups", "int_routing_barcode", "vrid", "origin", "sortation"):
        return (NA, "", "Specifico di altro vettore/cliente", "")
    return (NV, "", "Nessuna sorgente EURODIS", "")

wb = openpyxl.load_workbook(PATH)
src = wb["Tracciato unico preadvising"]
name_new = "Mapping EURODIS"
if name_new in wb.sheetnames:
    del wb[name_new]
ws = wb.create_sheet(name_new)
hdr = ["Path ICD", "Livello 1", "Livello 2", "Livello 3", "Livello 4", "Obbl. ICD", "Tipo", "Dim.",
       "Categoria", "Sorgente EURODIS", "Regola / note", "Punto aperto"]
ws.append(hdr)

stack = ["", "", "", ""]
used = set()
counts = {}
icd_info = {}
for row in src.iter_rows(min_row=2, values_only=True):
    lv = list(row[:4])
    if not any(lv):
        continue
    depth = next(i for i, v in enumerate(lv) if v)
    name = str(lv[depth]).strip()
    stack[depth] = name
    for j in range(depth + 1, 4):
        stack[j] = ""
    path = ".".join(stack[: depth + 1])
    is_container = (row[5] or "").lower().startswith(("object", "array")) or name == "importi"
    if is_container:
        rule = ("", "", "(contenitore)", "")
    else:
        rule = R.get(path) or default(path, name)
        used.add(path)
    counts[rule[0] or "contenitore"] = counts.get(rule[0] or "contenitore", 0) + 1
    icd_info[path] = (row[4], row[5], row[6])
    ws.append([path, *lv, row[4], row[5], row[6], *rule])

missing = sorted(set(R) - used)

# Campi EURODIS senza destinazione ICD
gaps = [
    ("truckLicensePlate", "Targa mezzo", "Nessun campo ICD; serve?", ""),
    ("lineHaulIndicator", "N/D/L", "Solo input per ce-rules", "PA-03"),
    ("entranceTimestamp", "Presa fisica", "Eventuale alternativa a data_spedizione", ""),
    ("referencedEurodisShipmentNumber", "Spedizione originaria (resi/rettifiche)", "Nessun campo ICD", "PA-17"),
    ("consignorsReference", "Riferimento mittente", "Nessun campo ICD libero (rif_spedizione e' edReference)", ""),
    ("consignorsCustomerNumberAtDispatchPartner", "Cliente presso partner", "Nessun campo ICD", ""),
    ("termsOfDeliveryText", "Incoterms", "Solo per addebito_oneri", "PA-13"),
    ("returnAddress", "Indirizzo di reso", "Nessun oggetto ICD per indirizzo di reso", "PA-17"),
    ("remunerationAddress", "Indirizzo di fatturazione", "Nessun oggetto ICD", ""),
    ("bookingInAddress", "Indirizzo di prenotazione consegna", "Nessun oggetto ICD", ""),
    ("consigneeAddress.alternativeContactTelephone", "Telefono alternativo", "Nessun campo ICD", ""),
    ("consigneeAddress.pudoID", "Consegna a punto di ritiro (PUDO)", "Nessun campo ICD generico (solo dhl_pudo*)", "PA-20"),
    ("consigneeAddress.gatewayAssigned", "Gateway assegnato", "Nessun campo ICD", ""),
    ("consigneePreAdvice.telephonePreAdvice", "Preavviso telefonico", "Nessun campo ICD", "PA-18"),
    ("deliveryService.secureDeliveryId", "Consegna sicura", "Nessun campo ICD", ""),
    ("deliveryService.deliveryBookInRequired", "Prenotazione consegna", "Forse servizio accessorio via ce-rules", "PA-03"),
    ("deliveryService.completeDeliveryOnly", "Consegna solo completa", "Nessun campo ICD", ""),
    ("deliveryService.additionalNotification", "Notifica a terzi", "Nessun campo ICD", ""),
    ("deliveryService.deliveryDateWindowEnd", "Fine finestra data", "ICD ha solo data_consegna", ""),
    ("cashOnDeliveryService.iban / bic / accountHolder", "Conto di accredito contrassegno", "ICD ha IBAN solo in cliente (da SFA)", "PA-19"),
    ("cashOnDeliveryService.paymentReason / paymentReference", "Causale contrassegno", "Nessun campo ICD", "PA-19"),
    ("cashOnDeliveryService.currency", "Valuta contrassegno", "ICD non ha valuta del contrassegno", "PA-19"),
    ("customsService.sign / number", "Riferimenti doganali", "Nessun campo ICD", "PA-13"),
    ("hazardousADRServices[]", "Merci pericolose ADR", "Nessun oggetto ICD", "PA-13"),
    ("hazardousCollis[]", "Colli ADR", "Nessun oggetto ICD", "PA-13"),
    ("collis[].referencedEuroplateBarcode", "Barcode collo originario (resi)", "Nessun campo ICD", "PA-17"),
    ("collis[].volumetric / volumetricWeight", "Volume e peso volumetrico", "Nessun campo ICD", ""),
    ("collis[].boxType / boxDimensionCategory / boxId", "Obsoleti secondo EURODIS", "Ignorare", ""),
]
name_gap = "EURODIS non mappati"
if name_gap in wb.sheetnames:
    del wb[name_gap]
wg = wb.create_sheet(name_gap)
wg.append(["Campo EURODIS", "Significato", "Motivo / note", "Punto aperto"])
for g in gaps:
    wg.append(list(g))

name_leg = "Legenda mapping"
if name_leg in wb.sheetnames:
    del wb[name_leg]
wl = wb.create_sheet(name_leg)
for r in [
    ["Categoria", "Significato"],
    [D, "Copiato dal campo EURODIS senza modifiche"],
    [T, "Dal campo EURODIS con una regola (formato, decodifica, troncamento)"],
    [C, "Valore fisso per tutto il canale EURODIS"],
    [G, "Prodotto dal batch (timestamp, id esecuzione)"],
    [F1, "Dati del cliente EURODIS da F1 (PA-10)"],
    [CE, "Derivato da ce-rules (PA-03)"],
    [AV, "Valorizzato a valle (normalizzazione, arricchimento, validazione): il batch lo lascia vuoto"],
    [NA, "Non pertinente al flusso EURODIS"],
    [DD, "Serve una decisione: vedi punto aperto"],
    [NV, "Nessuna sorgente EURODIS: il campo resta vuoto"],
    [],
    ["Fonte EURODIS", "https://testapi.eurodis.com/receive-pan/swagger/Release%20v1/swagger.json (TEST, letto 2026-10-22)"],
    ["Manutenzione", "Mary (analista), iniziativa 136596. Registro punti aperti in iniziative/136596/analisi.md"],
]:
    wl.append(r)

# --- Mapping inverso: un campo EURODIS per riga (o una riga per ogni destinazione ICD) ---
# Campi di GetShipmentResponse2 (swagger receive-pan Release v1): (path, tipo, max, obbligatorio)
ADDR = [("contactName", "string", 40, "NO"), ("contactTelephone", "string", 30, "NO"),
        ("alternativeContactTelephone", "string", 20, "NO"), ("contactSMS", "string", 35, "NO"),
        ("contactEMail", "string", 80, "NO"), ("companyName", "string", 40, "SI"),
        ("additionalAddressInfo1", "string", 40, "NO"), ("additionalAddressInfo2", "string", 40, "NO"),
        ("street", "string", 40, "SI"), ("streetNumber", "string", 40, "NO"), ("region", "string", 40, "NO"),
        ("city", "string", 35, "SI"), ("zipcode", "string", 9, "SI"), ("countryCode", "string", 2, "SI"),
        ("additionalConsigneeInfo", "string", 70, "NO"), ("gatewayAssigned", "string", 3, "NO")]
EF = [
    ("edReference", "number", "", "SI"), ("referencedEurodisShipmentNumber", "integer", "", "NO"),
    ("shipmentTypeIndicator", "string", 1, "SI"), ("lineHaulIndicator", "string", 1, "SI"),
    ("clearanceTimestamp", "date-time", "", "SI"), ("entranceTimestamp", "date-time", "", "NO"),
    ("numberOfHandlingUnits", "integer", "", "SI"), ("shipmentWeight", "number", "", "SI"),
    ("consignorsCustomerNumberAtDispatchPartner", "string", 35, "NO"), ("consignorsReference", "string", 35, "NO"),
    ("dispatchPartnersDomesticShipmentNumber", "string", 35, "SI"), ("truckLicensePlate", "string", 40, "NO"),
    ("termsOfDeliveryText", "string", 35, "NO"),
    *[("consigneeAddress." + n, t, l, o) for n, t, l, o in ADDR], ("consigneeAddress.pudoID", "string", 35, "NO"),
    *[("shipperAddress." + n, t, l, o) for n, t, l, o in ADDR],
    ("returnAddress", "object", "", "NO"), ("remunerationAddress", "object", "", "NO"), ("bookingInAddress", "object", "", "NO"),
    ("deliveryService.deliveryService", "string", 1, "SI"), ("deliveryService.deliveryType", "string", 1, "SI"),
    ("deliveryService.secureDeliveryId", "string", 40, "NO"), ("deliveryService.deliveryBookInRequired", "boolean", "", "NO"),
    ("deliveryService.deliveryTimeWindowStart", "string", 8, "NO"), ("deliveryService.deliveryTimeWindowEnd", "string", 8, "NO"),
    ("deliveryService.deliveryDateWindowStart", "string", "", "NO"), ("deliveryService.deliveryDateWindowEnd", "string", "", "NO"),
    ("deliveryService.completeDeliveryOnly", "boolean", "", "NO"), ("deliveryService.additionalNotification", "boolean", "", "NO"),
    ("cashOnDeliveryService", "object", "", "NO"),
    ("cashOnDeliveryService.iban", "string", 34, "NO"), ("cashOnDeliveryService.bic", "string", 11, "NO"),
    ("cashOnDeliveryService.accountHolder", "string", 35, "SI"), ("cashOnDeliveryService.paymentReason", "string", 140, "NO"),
    ("cashOnDeliveryService.paymentReference", "string", 35, "NO"), ("cashOnDeliveryService.amount", "number", "", "SI"),
    ("cashOnDeliveryService.currency", "string", 3, "SI"), ("cashOnDeliveryService.identifierForCODCollection", "integer", "", "SI"),
    ("customsService", "object", "", "NO"),
    ("customsService.commercialValueOfGoods", "number", "", "NO"), ("customsService.currency", "string", 3, "NO"),
    ("customsService.dutyCode", "string", 12, "SI"), ("customsService.sign", "string", 10, "NO"),
    ("customsService.number", "integer", "", "NO"), ("customsService.descriptionOfGoods", "string", 20, "SI"),
    ("hazardousADRServices[]", "array", "", "NO"), ("hazardousCollis[]", "array", "", "NO"),
    ("consigneePreAdvice.smsPreAdvice", "boolean", "", "SI"), ("consigneePreAdvice.eMailPreAdvice", "boolean", "", "SI"),
    ("consigneePreAdvice.telephonePreAdvice", "boolean", "", "SI"),
    ("collis[]", "array", "", "SI"),
    ("collis[].identifierPalletOrCollo", "string", 2, "SI"), ("collis[].weightOfUnit", "number", "", "SI"),
    ("collis[].europlateBarcode", "string", 24, "NO"), ("collis[].referencedEuroplateBarcode", "string", 24, "NO"),
    ("collis[].consignorsColloBarcode", "string", 50, "NO"), ("collis[].length", "integer", "", "SI"),
    ("collis[].width", "integer", "", "SI"), ("collis[].height", "integer", "", "SI"),
    ("collis[].volumetric", "integer", "", "NO"), ("collis[].volumetricWeight", "number", "", "NO"),
    ("collis[].boxType / boxDimensionCategory / boxId", "string", "", "NO"),
]
import re, json, os

# Descrizioni originali dallo swagger EURODIS salvato accanto a questo script
SWAGGER = os.path.join(os.path.dirname(PATH), "eurodis-receive-pan-swagger-v1.json")
_sc = json.load(open(SWAGGER, encoding="utf-8"))["components"]["schemas"]
SCHEMA_OF = {"": "GetShipmentResponse2", "consigneeAddress": "GetShipmentAddressPudoResponse",
             "shipperAddress": "GetShipmentAddressResponse", "deliveryService": "GetShipmentDeliveryServiceResponse",
             "cashOnDeliveryService": "GetShipmentCashOnDeliveryServiceResponse",
             "customsService": "GetShipmentCustomsServiceResponse",
             "consigneePreAdvice": "GetShipmentConsigneePreAdviceResponse", "collis[]": "GetShipmentColliDataResponse2"}

def desc_en(field):
    obj, _, leaf = field.rpartition(".")
    if " / " in leaf:
        leaf = leaf.split(" / ")[0]
    props = _sc[SCHEMA_OF.get(obj, "GetShipmentResponse2")].get("properties", {})
    d = (props.get(leaf) or {}).get("description") or ""
    return re.sub(r"\s*<br\s*/?>\s*", " ", d).strip()

# Traduzioni italiane. Chiave: campo EURODIS; per gli indirizzi "addr.<campo>" vale per tutti gli indirizzi.
IT = {
    "edReference": "Numero univoco della spedizione nella rete EURODIS. Per la definizione di dettaglio vedere il documento A02 del manuale di rete EURODIS.",
    "referencedEurodisShipmentNumber": "Numero EURODIS della spedizione originaria. Obbligatorio per i tipi \"R\", \"D\" e \"A\"; condizionato per \"P\"; facoltativo per \"E\".",
    "shipmentTypeIndicator": "N = normale; A = ordine di ritiro; P = spedizione risultante da un ordine di ritiro; R = reso; X = AUX-PAN; B = spedizione normale in neighbourhood business (NB); D = reso in NB; E = ritiro in NB; C = spedizione B2C (facoltativo, da concordare fra partner mittente e partner di consegna); Y = B2C in NB; L = richiesta di reso; M = reso dal destinatario (solo verso il partner di consegna iniziale; trasmesso come \"R\" ai partner di transito e consegna, come \"E\" in NB); W = spedizione distribuita senza EUROPLATE.",
    "lineHaulIndicator": "Deve essere \"N\" (normale), \"D\" (iniezione diretta) o \"L\" (trazione diretta).",
    "clearanceTimestamp": "Data e ora di creazione del dataset PAN della spedizione. Es. \"2000-01-31T18:30:00.000Z\" = 31/01/2000 ore 18:30.",
    "entranceTimestamp": "Data di ingresso della spedizione nel sistema domestico del partner, cioè quando è stata fisicamente consegnata. Es. \"2000-01-31T18:30:00.000Z\".",
    "numberOfHandlingUnits": "Totale di colli e pallet. Deve coincidere con il numero di colli trasmessi.",
    "shipmentWeight": "Peso in chilogrammi con due decimali.",
    "consignorsCustomerNumberAtDispatchPartner": "(Nessuna descrizione nello swagger) Codice cliente del mittente presso il partner che spedisce.",
    "consignorsReference": "(Nessuna descrizione nello swagger) Riferimento del mittente.",
    "dispatchPartnersDomesticShipmentNumber": "(Nessuna descrizione nello swagger) Numero di spedizione domestico del partner che spedisce.",
    "truckLicensePlate": "Targa del camion (e del rimorchio).",
    "termsOfDeliveryText": "Facoltativo per spedizioni intra-UE, obbligatorio per paesi extra-UE. Deve essere un Incoterm ufficiale DPU o EXW (per ritiri e resi) oppure uno degli Incoterm interni di rete DDC e DDX.",
    "addr.contactName": "Nome del contatto all'indirizzo, stampato sulla 2a riga indirizzo dell'EUROPLATE.",
    "consigneeAddress.contactName": "Nome del contatto all'indirizzo, stampato sulla 2a riga indirizzo dell'EUROPLATE. Per le spedizioni PUDO contiene il nome del negozio.",
    "addr.contactTelephone": "Condizione per il servizio di prenotazione consegna: se il servizio è concordato fra i partner e il flag di prenotazione è presente in DeliveryService.",
    "addr.alternativeContactTelephone": "Telefono di contatto alternativo (alfanumerico per ammettere il prefisso internazionale).",
    "addr.contactSMS": "Numero SMS del contatto principale. Formato: prefisso internazionale (+XX) seguito dal numero. Es. +491711234567.",
    "addr.contactEMail": "Indirizzo email del contatto principale; vedere anche l'allegato A09 - matrice del servizio di preavviso via email.",
    "addr.companyName": "Per spedizioni \"toB\" nome dell'azienda o organizzazione; per spedizioni \"toC\" nome del destinatario. Stampato sulla 1a riga indirizzo dell'EUROPLATE.",
    "addr.additionalAddressInfo1": "Informazioni aggiuntive sull'indirizzo (reparto, edificio, piano, codice porta, stanza, numero locker...). Stampato sulla 3a riga indirizzo dell'EUROPLATE.",
    "addr.additionalAddressInfo2": "Informazioni aggiuntive sull'indirizzo (reparto, edificio, piano, codice porta, stanza, numero locker...). Stampato sulla 4a riga indirizzo dell'EUROPLATE.",
    "consigneeAddress.additionalAddressInfo2": "Informazioni aggiuntive sull'indirizzo (reparto, edificio, piano, codice porta, stanza, numero locker...). Stampato sulla 4a riga indirizzo dell'EUROPLATE. Per le spedizioni PUDO contiene PS (Parcel Shop) o PT (Parcel Terminal).",
    "addr.street": "Via + numero civico, se il civico non è trasmesso a parte nel campo StreetNumber. Stampato sulla 5a riga indirizzo dell'EUROPLATE.",
    "addr.streetNumber": "Numero civico, stampato sulla 5a riga dell'EUROPLATE. Solo se il sistema del partner mittente supporta la trasmissione separata del civico; altrimenti il civico va nello stesso campo della via.",
    "addr.region": "Regione, provincia o contea.",
    "addr.city": "Città o località postale, stampata sulla 6a riga dell'EUROPLATE troncata a 30 caratteri.",
    "addr.zipcode": "CAP, stampato sulla 6a riga dell'EUROPLATE (alfanumerico per il Regno Unito, senza indicatore di nazione; formato validato, corrispondenza con l'indirizzo no).",
    "addr.countryCode": "Codice nazione ISO 3166 alpha-2, verificato sul database della piattaforma. Stampato sulla 6a riga dell'EUROPLATE.",
    "addr.additionalConsigneeInfo": "Informazioni aggiuntive del mittente per il destinatario o istruzioni di consegna.",
    "addr.gatewayAssigned": "Codice del gateway assegnato all'indirizzo (gateway di ingresso per l'indirizzo di consegna, di uscita per quello del mittente), come da manuale di rete A12.",
    "consigneeAddress.pudoID": "Id della stazione PUDO (punto di ritiro e consegna) per le spedizioni con servizio di consegna Z.",
    "returnAddress": "(Oggetto) Indirizzo di reso.",
    "remunerationAddress": "(Oggetto) Indirizzo di fatturazione/remunerazione.",
    "bookingInAddress": "(Oggetto) Indirizzo per la prenotazione della consegna.",
    "deliveryService.deliveryService": "Deve essere D (giorno), T (a orario), M (mattina), A (pomeriggio), E (sera), 6 (sabato), 7 (domenica), B (su prenotazione), W (finestra di consegna), F (entro le 10:00) o Z (consegna in punto PUDO).",
    "deliveryService.deliveryType": "Deve essere S (firma richiesta), N (firma non richiesta), D (deposito sicuro), I (identificazione sicura), B (consolidamento presso il partner di drop-off) o C (consolidamento presso il partner di consegna diverso dall'iniziatore). B e C valgono solo per resi via PUDO.",
    "deliveryService.secureDeliveryId": "Id del deposito sicuro o id da usare per la consegna sicura.",
    "deliveryService.deliveryBookInRequired": "Vero/falso: il partner di consegna deve concordare la prenotazione della consegna.",
    "deliveryService.deliveryTimeWindowStart": "Solo per il servizio a finestra di consegna e per servizi a orario specifico. Formato HH:mm:ss.",
    "deliveryService.deliveryTimeWindowEnd": "Solo per il servizio a finestra di consegna. Formato HH:mm:ss.",
    "deliveryService.deliveryDateWindowStart": "Solo per il servizio a finestra di consegna con data specifica. Formato yyyy-MM-dd.",
    "deliveryService.deliveryDateWindowEnd": "Solo per il servizio a finestra di consegna. Formato yyyy-MM-dd.",
    "deliveryService.completeDeliveryOnly": "Vero se il cliente vuole tutti i colli in un'unica consegna, altrimenti falso.",
    "deliveryService.additionalNotification": "Dati di contatto di un terzo da avvisare della consegna.",
    "cashOnDeliveryService": "(Oggetto) Contrassegno: presente se la spedizione è in contrassegno.",
    "cashOnDeliveryService.iban": "Obbligatorio se IdentifierForCODCollection = 3.",
    "cashOnDeliveryService.bic": "Obbligatorio se IdentifierForCODCollection = 3.",
    "cashOnDeliveryService.accountHolder": "Intestatario del conto.",
    "cashOnDeliveryService.paymentReason": "Attualmente, per il contrassegno fra Post AT e TOF-DE, si usano solo 13 cifre (numero spedizione EURODIS).",
    "cashOnDeliveryService.paymentReference": "Attualmente, per il contrassegno fra Post AT e TOF-DE, si usano solo 12 cifre (numero spedizione EURODIS).",
    "cashOnDeliveryService.amount": "Importo nella valuta indicata, con due decimali.",
    "cashOnDeliveryService.currency": "EUR o altro codice valuta ISO.",
    "cashOnDeliveryService.identifierForCODCollection": "Contrassegno incassato in contanti o con assegno; i contanti sono lo standard per Post AT e la maggior parte dei partner. Deve essere 3 (contanti) o 5 (assegno).",
    "customsService": "(Oggetto) Dati doganali.",
    "customsService.commercialValueOfGoods": "Per motivi assicurativi; obbligatorio se non si tratta di documenti.",
    "customsService.currency": "Codice valuta se il valore commerciale è valorizzato (EUR o altro codice ISO).",
    "customsService.dutyCode": "Obbligatorio per extra-UE (= codice della tariffa doganale).",
    "customsService.sign": "Descrizione aggiuntiva facoltativa della spedizione (numero di identificazione/riferimento).",
    "customsService.number": "Descrizione aggiuntiva facoltativa della spedizione (numero di identificazione/riferimento).",
    "customsService.descriptionOfGoods": "Informazione facoltativa sui camion di trasporto, es. per la Svizzera. NB: la descrizione dello swagger non corrisponde al nome del campo (\"descrizione merce\"), da chiarire con EURODIS.",
    "hazardousADRServices[]": "(Array) Dati delle merci pericolose ADR: classe di pericolo, numero ONU, gruppo di imballaggio, peso lordo ADR, codice galleria, categoria di trasporto, denominazione.",
    "hazardousCollis[]": "(Array) Colli ADR: numero di colli per tipo e descrizione ADR del tipo di collo.",
    "consigneePreAdvice.smsPreAdvice": "Se vero, ContactSMS in ConsigneeAddress deve essere valorizzato.",
    "consigneePreAdvice.eMailPreAdvice": "Se vero, ContactEMail in ConsigneeAddress deve essere valorizzato.",
    "consigneePreAdvice.telephonePreAdvice": "Se vero, ContactTelephone in ConsigneeAddress deve essere valorizzato.",
    "collis[]": "(Array) Dati dei colli: un elemento per collo o pallet.",
    "collis[].identifierPalletOrCollo": "\"P\" (pallet), \"C\" (collo), \"P+\" (pallet+) o \"C+\" (collo+).",
    "collis[].weightOfUnit": "In chilogrammi.",
    "collis[].europlateBarcode": "Barcode EURODIS del collo; valore di default (0) in caso di neighbourhood business.",
    "collis[].referencedEuroplateBarcode": "Valorizzato solo se il collo fa parte di un reso: riporta il barcode EURODIS del collo originario.",
    "collis[].consignorsColloBarcode": "Barcode domestico di reso del partner di consegna originario, generato/assegnato dalla soluzione resi EURODIS, oppure barcode del cliente o del partner mittente. Per il neighbourhood business va fornito il barcode domestico del partner secondo la SOP Neighbourhood business.",
    "collis[].length": "In centimetri.",
    "collis[].width": "In centimetri.",
    "collis[].height": "In centimetri.",
    "collis[].volumetric": "Metri cubi.",
    "collis[].volumetricWeight": "In chilogrammi.",
    "collis[].boxType / boxDimensionCategory / boxId": "Forse obsoleti, da discutere.",
}

def desc_it(field):
    if field in IT:
        return IT[field]
    obj, _, leaf = field.rpartition(".")
    return IT.get("addr." + leaf, "") if obj in ("consigneeAddress", "shipperAddress") else ""

def descrizione(field):
    en, it = desc_en(field), desc_it(field)
    return f"{en}\n\nIT: {it}" if en else f"IT: {it}"

def tokens(source):
    return [t for t in re.split(r"[\s,/?()]+", source or "") if t]

def matches(field, tok):
    if tok.endswith(".*"):
        return field.startswith(tok[:-1])
    return field == tok or field.endswith("." + tok)

def gap_keys(key):
    parts = [p.strip() for p in key.split(" / ")]
    prefix = parts[0].rsplit(".", 1)[0] + "." if "." in parts[0] else ""
    return [parts[0]] + [p if "." in p or not prefix else prefix + p for p in parts[1:]]

gap_by_field = {}
for g in gaps:
    for k in gap_keys(g[0]):
        gap_by_field[k] = g
gap_by_field["collis[].boxType / boxDimensionCategory / boxId"] = gaps[-1]
gap_by_field.update(GAP_23)

# Commenti scritti a mano nella colonna A del foglio: si conservano a ogni rigenerazione.
# Chiave (campo EURODIS, campo ICD); se la coppia non esiste piu' il commento va sulla
# prima riga del campo EURODIS, con l'indicazione della riga originaria.
old_comments = {}
if "Mapping inverso EURODIS" in wb.sheetnames:
    _old = wb["Mapping inverso EURODIS"]
    _h = [c.value for c in _old[1]]
    if _h and _h[0] == "Commenti":
        _ie, _ii = _h.index("Campo EURODIS"), _h.index("Campo ICD")
        for r in _old.iter_rows(min_row=2, values_only=True):
            if r[0]:
                old_comments[(r[_ie], r[_ii])] = r[0]

name_inv = "Mapping inverso EURODIS"
if name_inv in wb.sheetnames:
    del wb[name_inv]
wi = wb.create_sheet(name_inv, index=1)
wi.append(["Commenti", "Campo EURODIS", "Descrizione EURODIS (originale / IT)", "Tipo", "Max", "Obbl. EURODIS",
           "Campo ICD", "Obbl. ICD", "Dim. ICD", "Categoria", "Regola / note", "Punto aperto"])
inv_counts = {"mappati": 0, "non mappati": 0}
no_it = []
for field, tp, ln, ob in EF:
    if not desc_it(field):
        no_it.append(field)
    ds = descrizione(field)
    targets = [(p, r) for p, r in R.items() if any(matches(field, t) for t in tokens(r[1]))]
    if targets:
        inv_counts["mappati"] += 1
        for p, r in sorted(targets):
            o, _, dim = icd_info.get(p, ("", "", ""))
            wi.append([old_comments.pop((field, p), None), field, ds, tp, ln, ob, p, o, dim, r[0], r[2], r[3]])
    else:
        inv_counts["non mappati"] += 1
        g = gap_by_field.get(field)
        wi.append([old_comments.pop((field, "—"), None), field, ds, tp, ln, ob, "—", "", "", "NON MAPPATO",
                   g[2] if g else "Nessun campo ICD", g[3] if g else ""])
# commenti rimasti orfani (la coppia EURODIS -> ICD non esiste piu')
for (field, icd), txt in old_comments.items():
    for r in wi.iter_rows(min_row=2):
        if r[1].value == field:
            r[0].value = ((r[0].value + "\n") if r[0].value else "") + f"[era su {icd}] {txt}"
            break
fills_inv = {"NON MAPPATO": "F8CBAD"}

fills = {D: "C6EFCE", T: "DDEBF7", C: "E2EFDA", G: "E2EFDA", F1: "FFF2CC", CE: "FFF2CC",
         AV: "EDEDED", NA: "EDEDED", DD: "F8CBAD", NV: "FFFFFF"}
for r in wi.iter_rows(min_row=2):
    cat = r[9].value
    color = fills.get(cat) or fills_inv.get(cat)
    if color:
        r[9].fill = PatternFill("solid", fgColor=color)
    for c in r:
        c.alignment = Alignment(wrap_text=c.column in (1, 3, 11), vertical="top")
    if r[0].value:
        r[0].fill = PatternFill("solid", fgColor="FFF2CC")
for sheet, widths in ((ws, [45, 14, 14, 14, 14, 9, 10, 7, 16, 45, 70, 12]),
                      (wi, [40, 44, 70, 10, 6, 9, 42, 9, 8, 16, 60, 12]),
                      (wg, [45, 40, 60, 12]), (wl, [18, 100])):
    for c in sheet[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for i, w in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
for r in ws.iter_rows(min_row=2):
    cat = r[8].value
    if cat in fills:
        r[8].fill = PatternFill("solid", fgColor=fills[cat])
    if cat == "":
        for c in r:
            c.font = Font(bold=True)
    r[10].alignment = Alignment(wrap_text=True, vertical="top")
for r in wl.iter_rows(min_row=2):
    if r[0].value in fills:
        r[0].fill = PatternFill("solid", fgColor=fills[r[0].value])

wb.save(PATH)
print("righe:", ws.max_row - 1, counts)
print("regole non agganciate:", missing)
print("mapping inverso:", inv_counts, "righe:", wi.max_row - 1)
print("campi senza traduzione:", no_it)
